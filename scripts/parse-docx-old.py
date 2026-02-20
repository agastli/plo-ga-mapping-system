#!/usr/bin/env python3
"""
Word Document Parser for PLO-GA Mapping System
Extracts program information, PLOs, mapping matrix, and justifications from .docx files
"""

import sys
import json
import re
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

def extract_program_info(doc):
    """Extract program name, department, and college from document"""
    program_info = {
        "programNameEn": "",
        "programNameAr": "",
        "departmentEn": "",
        "departmentAr": "",
        "collegeEn": "",
        "collegeAr": "",
        "language": "en"
    }
    
    # Look for program information in the first few paragraphs
    for para in doc.paragraphs[:20]:
        text = para.text.strip()
        
        # Check for Arabic content to determine language
        if re.search(r'[\u0600-\u06FF]', text):
            program_info["language"] = "ar"
        
        # Extract program name (usually after "Program:" or "البرنامج:")
        if "Program:" in text or "البرنامج:" in text:
            program_info["programNameEn"] = text.split("Program:")[-1].strip() if "Program:" in text else ""
            program_info["programNameAr"] = text.split("البرنامج:")[-1].strip() if "البرنامج:" in text else ""
        
        # Extract department
        if "Department:" in text or "القسم:" in text:
            program_info["departmentEn"] = text.split("Department:")[-1].strip() if "Department:" in text else ""
            program_info["departmentAr"] = text.split("القسم:")[-1].strip() if "القسم:" in text else ""
        
        # Extract college
        if "College:" in text or "الكلية:" in text:
            program_info["collegeEn"] = text.split("College:")[-1].strip() if "College:" in text else ""
            program_info["collegeAr"] = text.split("الكلية:")[-1].strip() if "الكلية:" in text else ""
    
    return program_info

def extract_plos(doc):
    """Extract Program Learning Outcomes from document"""
    plos = []
    in_plo_section = False
    plo_pattern = re.compile(r'^(PLO\s*\d+)[:\s]*(.+)', re.IGNORECASE)
    
    for para in doc.paragraphs:
        text = para.text.strip()
        
        # Start capturing PLOs after finding the PLO section header
        if "Program Learning Outcomes" in text or "مخرجات التعلم" in text:
            in_plo_section = True
            continue
        
        # Stop at mapping section
        if "Mapping" in text or "المواءمة" in text or in_plo_section and len(plos) > 0 and not text.startswith("PLO"):
            break
        
        if in_plo_section:
            match = plo_pattern.match(text)
            if match:
                code = match.group(1).replace(" ", "").upper()
                description = match.group(2).strip()
                
                # Determine if it's English or Arabic
                is_arabic = bool(re.search(r'[\u0600-\u06FF]', description))
                
                plos.append({
                    "code": code,
                    "descriptionEn": description if not is_arabic else "",
                    "descriptionAr": description if is_arabic else "",
                    "sortOrder": len(plos) + 1
                })
    
    return plos

def extract_mapping_matrix(doc):
    """Extract the PLO-to-Competency mapping matrix from tables"""
    mappings = {}
    
    # Find the mapping table (usually the largest table)
    for table in doc.tables:
        if table.rows and len(table.rows) > 20:  # Mapping table should have 21+ rows for competencies
            # Parse header row to get PLO columns
            header_row = table.rows[0]
            plo_columns = []
            
            for i, cell in enumerate(header_row.cells):
                cell_text = cell.text.strip()
                if re.match(r'PLO\s*\d+', cell_text, re.IGNORECASE):
                    plo_code = re.sub(r'\s+', '', cell_text).upper()
                    plo_columns.append((i, plo_code))
            
            # Parse data rows
            for row in table.rows[1:]:
                cells = row.cells
                if len(cells) < 2:
                    continue
                
                # First column should contain competency code (C1-1, C1-2, etc.)
                competency_code = ""
                for cell in cells[:3]:  # Check first few cells for competency code
                    text = cell.text.strip()
                    match = re.search(r'(C\d+-\d+)', text)
                    if match:
                        competency_code = match.group(1)
                        break
                
                if not competency_code:
                    continue
                
                # Extract weights for each PLO
                for col_idx, plo_code in plo_columns:
                    if col_idx < len(cells):
                        weight_text = cells[col_idx].text.strip()
                        try:
                            weight = float(weight_text)
                            if 0.0 <= weight <= 1.0:
                                key = f"{plo_code}_{competency_code}"
                                mappings[key] = {
                                    "ploCode": plo_code,
                                    "competencyCode": competency_code,
                                    "weight": weight
                                }
                        except ValueError:
                            # Not a valid number, skip
                            pass
    
    return list(mappings.values())

def extract_justifications(doc):
    """Extract justifications for each competency"""
    justifications = []
    in_justification_section = False
    current_ga = None
    
    # GA patterns
    ga_patterns = [
        (r'Graduate Attribute 1[:\s]*Competent', 'GA1'),
        (r'Graduate Attribute 2[:\s]*Life-long Learner', 'GA2'),
        (r'Graduate Attribute 3[:\s]*Well Rounded', 'GA3'),
        (r'Graduate Attribute 4[:\s]*Ethically.*Responsible', 'GA4'),
        (r'Graduate Attribute 5[:\s]*Entrepreneurial', 'GA5'),
        (r'السمة الأولى[:\s]*الكفاءة', 'GA1'),
        (r'السمة الثانية[:\s]*المتعلم مدى الحياة', 'GA2'),
        (r'السمة الثالثة[:\s]*المتكامل', 'GA3'),
        (r'السمة الرابعة[:\s]*المسؤول', 'GA4'),
        (r'السمة الخامسة[:\s]*الريادي', 'GA5'),
    ]
    
    # Competency pattern (C1-1, C1-2, etc.)
    competency_pattern = re.compile(r'^(C\d+-\d+)\s*[–-]\s*(.+?):\s*(.+)', re.DOTALL)
    
    for para in doc.paragraphs:
        text = para.text.strip()
        
        # Start capturing after "Justifications" header
        if "Justifications for Mapping" in text or "مبررات المواءمة" in text:
            in_justification_section = True
            continue
        
        if not in_justification_section:
            continue
        
        # Check if this is a GA header
        for pattern, ga_code in ga_patterns:
            if re.search(pattern, text, re.IGNORECASE):
                current_ga = ga_code
                break
        
        # Check if this paragraph contains a competency justification
        if current_ga and text:
            match = competency_pattern.match(text)
            if match:
                competency_code = match.group(1)
                competency_name = match.group(2).strip()
                justification_text = match.group(3).strip()
                
                is_arabic = bool(re.search(r'[\u0600-\u06FF]', justification_text))
                
                justifications.append({
                    "gaCode": current_ga,
                    "competencyCode": competency_code,
                    "competencyName": competency_name,
                    "textEn": justification_text if not is_arabic else "",
                    "textAr": justification_text if is_arabic else ""
                })
    
    return justifications

def parse_document(file_path):
    """Main function to parse the Word document"""
    try:
        doc = Document(file_path)
        
        result = {
            "success": True,
            "data": {
                "programInfo": extract_program_info(doc),
                "plos": extract_plos(doc),
                "mappings": extract_mapping_matrix(doc),
                "justifications": extract_justifications(doc)
            }
        }
        
        return result
    
    except Exception as e:
        return {
            "success": False,
            "error": str(e)
        }

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"success": False, "error": "No file path provided"}))
        sys.exit(1)
    
    file_path = sys.argv[1]
    result = parse_document(file_path)
    print(json.dumps(result, ensure_ascii=False, indent=2))
